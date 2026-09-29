"""Pruebas de los generadores en Python y de la estructura del libro de análisis.

Ejecutar desde la raíz del proyecto:
    python -m unittest discover -s pruebas/unitarias -p "test_*.py" -v
"""

from __future__ import annotations

import base64
import io
import json
import re
import struct
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(RAIZ / "src" / "4-infraestructura" / "excel"))
sys.path.insert(0, str(RAIZ / "pruebas" / "aceptacion"))

import generar_datos_prueba  # noqa: E402
from generar_plantilla_comprobante import crear_plantilla, leer_contrato  # noqa: E402

LIBRO_DEMO = RAIZ / "salida" / "Analisis_Legalizaciones_TC_DEMO.xlsx"


def leer_consultas_power_query(ruta: Path) -> str:
    """Devuelve el código M (Section1.m) guardado dentro de un libro de Excel (formato MS-QDEFF)."""
    with zipfile.ZipFile(ruta) as libro:
        for nombre in libro.namelist():
            if not re.match(r"customXml/item\d+\.xml$", nombre):
                continue
            datos = libro.read(nombre)
            texto = datos.decode("utf-16") if datos[:2] in (b"\xff\xfe", b"\xfe\xff") else datos.decode("utf-8", "ignore")
            coincidencia = re.search(r"<DataMashup[^>]*>(.*?)</DataMashup>", texto, re.S)
            if not coincidencia:
                continue
            binario = base64.b64decode(coincidencia.group(1))
            largo = struct.unpack("<I", binario[4:8])[0]
            with zipfile.ZipFile(io.BytesIO(binario[8 : 8 + largo])) as paquete:
                return paquete.read("Formulas/Section1.m").decode("utf-8-sig")
    raise AssertionError(f"{ruta.name} no tiene consultas de Power Query")


class PlantillaComprobanteTest(unittest.TestCase):
    """La plantilla que llenan los colaboradores respeta su contrato (plantilla-comprobante.json)."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.contrato = leer_contrato()
        cls.carpeta = tempfile.TemporaryDirectory()
        cls.ruta = Path(cls.carpeta.name) / "plantilla.xlsx"
        crear_plantilla(cls.contrato).save(cls.ruta)
        cls.libro = load_workbook(cls.ruta)
        cls.hoja = cls.libro[cls.contrato["hoja"]]

    @classmethod
    def tearDownClass(cls) -> None:
        cls.libro.close()
        cls.carpeta.cleanup()

    def test_tabla_con_nombre_y_columnas_del_contrato(self) -> None:
        tabla = self.hoja.tables[self.contrato["tabla"]]
        fila = int(self.contrato["filaEncabezado"])
        self.assertTrue(tabla.ref.startswith(f"A{fila}:"))
        encabezados = [self.hoja.cell(row=fila, column=i + 1).value for i in range(len(self.contrato["columnas"]))]
        self.assertEqual(encabezados, [c["nombre"] for c in self.contrato["columnas"]])

    def test_filas_preparadas(self) -> None:
        tabla = self.hoja.tables[self.contrato["tabla"]]
        ultima_fila = int(re.search(r"(\d+)$", tabla.ref).group(1))
        self.assertEqual(ultima_fila - int(self.contrato["filaEncabezado"]), int(self.contrato["filasPreparadas"]))

    def test_validaciones_de_fecha_numero_y_lista(self) -> None:
        tipos = {v.type for v in self.hoja.data_validations.dataValidation}
        self.assertTrue({"date", "decimal", "list"} <= tipos)
        lista = next(v for v in self.hoja.data_validations.dataValidation if v.type == "list")
        self.assertIn("MXN", lista.formula1)

    def test_nombres_definidos_del_encabezado(self) -> None:
        for campo in self.contrato["encabezado"]["campos"]:
            self.assertIn(campo["nombreDefinido"], self.libro.defined_names)


class DatosPruebaTest(unittest.TestCase):
    """Los adjuntos de prueba son archivos válidos y cubren todos los escenarios."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.carpeta = tempfile.TemporaryDirectory()
        cls.destino = Path(cls.carpeta.name)
        generar_datos_prueba.generar(cls.destino)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.carpeta.cleanup()

    def test_pdfs_validos(self) -> None:
        for pdf in self.destino.glob("*.pdf"):
            datos = pdf.read_bytes()
            self.assertTrue(datos.startswith(b"%PDF-1.4"), pdf.name)
            inicio_xref = int(re.search(rb"startxref\n(\d+)", datos).group(1))
            self.assertTrue(datos[inicio_xref:].startswith(b"xref"), pdf.name)

    def test_xml_con_estructura_cfdi_4(self) -> None:
        raiz = ET.parse(self.destino / "cfdi_A001.xml").getroot()
        self.assertEqual(raiz.tag, "{http://www.sat.gob.mx/cfd/4}Comprobante")
        self.assertEqual(raiz.attrib["Version"], "4.0")

    def test_png_valido(self) -> None:
        self.assertTrue((self.destino / "factura_foto.png").read_bytes().startswith(b"\x89PNG\r\n\x1a\n"))

    def test_comprobante_con_plantilla_tiene_tres_lineas(self) -> None:
        libro = load_workbook(self.destino / "Gastos_Quincena.xlsx")
        hoja = libro["Gastos"]
        self.assertIn("tblGastos", hoja.tables)
        fila = int(leer_contrato()["filaEncabezado"])
        lineas = [f for f in hoja.iter_rows(min_row=fila + 1, values_only=True) if any(v is not None for v in f)]
        self.assertEqual(len(lineas), 3)
        libro.close()

    def test_comprobante_formato_libre_no_tiene_la_tabla(self) -> None:
        libro = load_workbook(self.destino / "Gastos_Formato_Libre.xlsx")
        self.assertFalse(any(hoja.tables for hoja in libro.worksheets))
        libro.close()

    def test_existen_los_adjuntos_de_todos_los_escenarios(self) -> None:
        escenarios = json.loads((RAIZ / "pruebas" / "aceptacion" / "Escenarios.json").read_text(encoding="utf-8"))["escenarios"]
        for escenario in escenarios:
            adjuntos = list(escenario.get("adjuntos", []))
            if "volumen" in escenario:
                volumen = escenario["volumen"]
                adjuntos += [a.replace("{nn}", f"{n:02d}") for n in range(1, volumen["cantidad"] + 1) for a in volumen["adjuntos"]]
            for adjunto in adjuntos:
                self.assertTrue((self.destino / adjunto).exists(), f"{escenario['id']}: falta {adjunto}")


@unittest.skipUnless(LIBRO_DEMO.exists(), "Genere antes el libro DEMO: Generar-LibroAnalisis.ps1 -Modo Demostracion")
class LibroAnalisisDemoTest(unittest.TestCase):
    """El libro de análisis generado trae sus consultas, tablas y validaciones."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.codigo_m = leer_consultas_power_query(LIBRO_DEMO)
        cls.libro = load_workbook(LIBRO_DEMO)

    @classmethod
    def tearDownClass(cls) -> None:
        cls.libro.close()

    def test_tiene_las_consultas(self) -> None:
        for consulta in ("pSitioUrl", "pModo", "Dominio", "Registro", "SoportesBase", "Gastos", "Gastos_Errores",
                         "Control_Operam", "Consolidado_Operam", "Alertas", "Resumen_Mensual", "fnLeerComprobante", "fnUtcALocal"):
            self.assertRegex(self.codigo_m, rf"shared {consulta} =", consulta)

    def test_modo_demostracion_y_estados_del_dominio(self) -> None:
        self.assertRegex(self.codigo_m, r'shared pModo = "Demostracion" meta')
        dominio = json.loads((RAIZ / "src" / "1-dominio" / "dominio.json").read_text(encoding="utf-8"))
        estados = ", ".join(f'"{e}"' for e in dominio["caso"]["estados"])
        self.assertIn("{" + estados + "}", self.codigo_m)

    def test_tablas_por_hoja(self) -> None:
        esperadas = {
            "Registro": "tblRegistro", "Alertas": "tblAlertas", "Consolidado_Operam": "tblConsolidadoOperam",
            "Control_Operam": "tblControlOperam", "Gastos": "tblGastosConsolidados", "Soportes": "tblSoportes",
            "Errores_Lectura": "tblErroresLectura", "Resumen": "tblResumenMensual",
        }
        for hoja, tabla in esperadas.items():
            self.assertIn(tabla, self.libro[hoja].tables, hoja)

    def test_lista_desplegable_de_estado(self) -> None:
        hoja = self.libro["Registro"]
        listas = [v for v in hoja.data_validations.dataValidation if v.type == "list"]
        self.assertTrue(any("lstEstados" in (v.formula1 or "") for v in listas))


if __name__ == "__main__":
    unittest.main()
