"""Genera la plantilla del comprobante quincenal que envían los colaboradores.

Capa de INFRAESTRUCTURA (Excel). El contrato de la plantilla (hoja, tabla, columnas,
validaciones) está en src/3-adaptadores/excel/plantilla-comprobante.json: la consulta
fnLeerComprobante de Power Query lee exactamente esa tabla, por eso las dos piezas
se generan desde el mismo archivo.

Uso:
    python src/4-infraestructura/excel/generar_plantilla_comprobante.py [--salida RUTA]

Requiere: openpyxl (lo instala Instalar-Prerrequisitos.ps1).
"""

from __future__ import annotations

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parents[3]
CONTRATO = RAIZ / "src" / "3-adaptadores" / "excel" / "plantilla-comprobante.json"
FECHA_MINIMA = date(2020, 1, 1)


def leer_contrato(ruta: Path = CONTRATO) -> dict:
    """Lee el contrato de la plantilla."""
    return json.loads(ruta.read_text(encoding="utf-8"))


def serial_excel(fecha: date) -> int:
    """Número de serie de Excel de una fecha (base 1899-12-30)."""
    return (fecha - date(1899, 12, 30)).days


def crear_plantilla(contrato: dict, filas_ejemplo: list[list] | None = None, colaborador: str = "", quincena: str = "") -> Workbook:
    """Crea el libro de la plantilla. Si se pasan filas de ejemplo, las escribe en la tabla."""
    libro = Workbook()
    hoja = libro.active
    hoja.title = contrato["hoja"]
    encabezado = contrato["encabezado"]
    fila_tabla = int(contrato["filaEncabezado"])
    columnas = contrato["columnas"]
    ultima_columna = get_column_letter(len(columnas))

    hoja["A1"] = encabezado["titulo"]
    hoja["A1"].font = Font(size=14, bold=True)

    fila = 3
    for campo in encabezado["campos"]:
        hoja.cell(row=fila, column=1, value=campo["etiqueta"]).font = Font(bold=True)
        celda = hoja.cell(row=fila, column=2)
        celda.fill = PatternFill("solid", fgColor="FFF2CC")
        libro.defined_names[campo["nombreDefinido"]] = DefinedName(campo["nombreDefinido"], attr_text=f"'{hoja.title}'!$B${fila}")
        fila += 1
    hoja["B3"] = colaborador or None
    hoja["B4"] = quincena or None

    hoja.cell(row=fila_tabla - 2, column=1, value=encabezado["instrucciones"]).font = Font(italic=True, color="595959")

    for indice, columna in enumerate(columnas, start=1):
        celda = hoja.cell(row=fila_tabla, column=indice, value=columna["nombre"])
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        hoja.column_dimensions[get_column_letter(indice)].width = columna.get("ancho", 14)

    filas = filas_ejemplo or []
    total_filas = max(int(contrato["filasPreparadas"]), len(filas))
    for desplazamiento, valores in enumerate(filas, start=1):
        for indice, valor in enumerate(valores, start=1):
            hoja.cell(row=fila_tabla + desplazamiento, column=indice, value=valor)

    ultima_fila = fila_tabla + total_filas
    tabla = Table(displayName=contrato["tabla"], ref=f"A{fila_tabla}:{ultima_columna}{ultima_fila}")
    tabla.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    hoja.add_table(tabla)

    # Formatos y validaciones por tipo de columna
    for indice, columna in enumerate(columnas, start=1):
        letra = get_column_letter(indice)
        rango = f"{letra}{fila_tabla + 1}:{letra}{ultima_fila}"
        tipo = columna["tipo"]
        if tipo == "fecha":
            validacion = DataValidation(type="date", operator="greaterThanOrEqual", formula1=str(serial_excel(FECHA_MINIMA)), allow_blank=True)
            validacion.error = "Escriba una fecha válida (dd/mm/aaaa)."
            formato = "dd/mm/yyyy"
        elif tipo == "numero":
            validacion = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            validacion.error = "Escriba un número mayor o igual a 0."
            formato = "#,##0.00"
        elif tipo == "lista":
            opciones = ",".join(columna["opciones"])
            validacion = DataValidation(type="list", formula1=f'"{opciones}"', allow_blank=True)
            validacion.error = f"Use uno de estos valores: {opciones}."
            formato = None
        else:
            validacion = None
            formato = "@"
        if validacion is not None:
            validacion.errorTitle = columna["nombre"]
            validacion.showErrorMessage = True
        else:
            validacion = DataValidation(type=None, allow_blank=True)
        validacion.promptTitle = columna["nombre"][:32]
        validacion.prompt = columna["ayuda"][:255]
        validacion.showInputMessage = True
        hoja.add_data_validation(validacion)
        validacion.add(rango)
        if formato:
            for (celda,) in hoja[rango]:
                celda.number_format = formato

    hoja.freeze_panes = f"A{fila_tabla + 1}"
    libro.properties.title = encabezado["titulo"]
    libro.properties.creator = "Legalizaciones TC - Emobility"
    return libro


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    contrato = leer_contrato()
    parser.add_argument("--salida", type=Path, default=RAIZ / "salida" / contrato["nombreArchivo"], help="Archivo .xlsx a generar")
    argumentos = parser.parse_args()
    argumentos.salida.parent.mkdir(parents=True, exist_ok=True)
    crear_plantilla(contrato).save(argumentos.salida)
    print(f"Plantilla generada: {argumentos.salida}")


if __name__ == "__main__":
    main()
